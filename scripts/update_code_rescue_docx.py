import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_box(doc, code_text, fill_hex="F8FAFC", border_hex="CBD5E1"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    lines = code_text.strip().split('\n')
    for i, line in enumerate(lines):
        line_num = f"{i+1:2d} | "
        r_num = p.add_run(line_num)
        r_num.font.name = 'Consolas'
        r_num.font.size = Pt(8)
        r_num.font.color.rgb = RGBColor(148, 163, 184)
        
        r_code = p.add_run(line + '\n' if i < len(lines)-1 else line)
        r_code.font.name = 'Consolas'
        r_code.font.size = Pt(8)
        r_code.font.color.rgb = RGBColor(15, 23, 42)
        
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def generate_docx(output_path):
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("Techastra '26 — Code Rescue\n")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p_title.add_run("OFFICIAL DEBUGGING QUESTION BANK (CALIBRATED & STREAMLINED)\n")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(2, 132, 199)
    
    p_desc = doc.add_paragraph("Tournament Blueprint: Round 1 (10Q, 2 bugs each, 15m) | Round 2 (5Q, 5 bugs each, 20m) | Round 3 (1Q, 10 bugs, 25m) — Total 35 Marks")
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc.runs[0].font.size = Pt(9)
    p_desc.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    p_desc.paragraph_format.space_after = Pt(12)
    
    # Table of Tournament Rules
    t = doc.add_table(rows=5, cols=8)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Round", "Stage Title", "Questions", "Bugs/Q", "Marks/Q", "Total Marks", "Duration", "Cutoff"]
    for i, h in enumerate(headers):
        cell = t.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(3, 105, 161)
        set_cell_background(cell, "E0F2FE")
        
    data = [
        ["Round 1", "Bug Hunt (Fast & Easy)", "10 Questions", "Exactly 2", "1 Mark", "10 Marks", "15 Mins", "≥ 5 Marks"],
        ["Round 2", "Logic Breaker (Liberal)", "5 Questions", "Exactly 5", "4 Marks", "20 Marks", "20 Mins", "≥ 10 Marks"],
        ["Round 3", "Code Rescue Finale", "1 Grand System", "Exactly 10", "5 Marks", "5 Marks", "25 Mins", "Leaderboard"],
        ["TOTAL", "All 3 Rounds Combined", "16 Challenges", "55 Total Bugs", "—", "35 Marks", "60 Mins", "Championship"]
    ]
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, text in enumerate(row_data):
            cell = t.cell(row_idx, col_idx)
            cell.text = text
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            if row_idx == 4:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_background(cell, "F8FAFC")
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ------------------ ROUND 1 ------------------
    h1 = doc.add_heading("ROUND 1 — BUG HUNT (10 Questions | 2 Bugs Each | 15 Mins)", level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(6)
    
    from competition_questions_data import r1_questions, r2_questions, r3_streamlined_buggy, r3_bugs_list, r3_streamlined_fixed
    
    for q in r1_questions:
        hq = doc.add_heading(f"{q['num']}. {q['title']} ({q['concept']})", level=2)
        hq.paragraph_format.space_before = Pt(8)
        hq.paragraph_format.space_after = Pt(2)
        
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"Work Order: {q['desc']}\n")
        r.font.size = Pt(8.5)
        r2 = p.add_run(f"Signature: {q['contract']} | Baseline: Input {q['sample_in']} -> Output: {q['sample_out']}")
        r2.font.size = Pt(8.5)
        r2.font.bold = True
        
        doc.add_paragraph("Buggy Starter Code (2 Errors Injected):").paragraph_format.space_after = Pt(1)
        add_code_box(doc, q['buggy_code'], fill_hex="FFF1F2", border_hex="FECDD3")
        
        doc.add_paragraph("Identified Errors:").paragraph_format.space_after = Pt(1)
        for b in q['bugs']:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(1)
            # strip html tags for docx
            clean_b = b.replace('<b>', '').replace('</b>', '').replace('<code>', '').replace('</code>', '')
            bp.add_run(clean_b).font.size = Pt(8)
            
        doc.add_paragraph("Verified Correct Solution:").paragraph_format.space_after = Pt(1)
        add_code_box(doc, q['fixed_code'], fill_hex="F0FDF4", border_hex="BBF7D0")
        
    # ------------------ ROUND 2 ------------------
    doc.add_page_break()
    h2 = doc.add_heading("ROUND 2 — LOGIC BREAKER (5 Questions | 5 Bugs Each | 20 Mins)", level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    
    for q in r2_questions:
        hq = doc.add_heading(f"{q['num']}. {q['title']} ({q['domain']})", level=2)
        hq.paragraph_format.space_before = Pt(8)
        hq.paragraph_format.space_after = Pt(2)
        
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"Work Order: {q['desc']}\n")
        r.font.size = Pt(8.5)
        r2 = p.add_run(f"Signature: {q['contract']} | Baseline: Input {q['sample_in']} -> Output: {q['sample_out']}")
        r2.font.size = Pt(8.5)
        r2.font.bold = True
        
        doc.add_paragraph("Streamlined Buggy Starter Code (5 Errors Injected):").paragraph_format.space_after = Pt(1)
        add_code_box(doc, q['buggy_code'], fill_hex="FFF1F2", border_hex="FECDD3")
        
        doc.add_paragraph("Identified Errors:").paragraph_format.space_after = Pt(1)
        for b in q['bugs']:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(1)
            clean_b = b.replace('<b>', '').replace('</b>', '').replace('<code>', '').replace('</code>', '')
            bp.add_run(clean_b).font.size = Pt(8)
            
        doc.add_paragraph("Verified Correct Solution:").paragraph_format.space_after = Pt(1)
        add_code_box(doc, q['fixed_code'], fill_hex="F0FDF4", border_hex="BBF7D0")

    # ------------------ ROUND 3 ------------------
    doc.add_page_break()
    h3 = doc.add_heading("ROUND 3 — CODE RESCUE FINALE (1 Grand System | 10 Bugs | 25 Mins)", level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph("Smart Restaurant Billing & Order Dispatch System (Streamlined to 38 Lines of Code)")
    p.runs[0].font.bold = True
    p.runs[0].font.size = Pt(10)
    
    doc.add_paragraph("Baseline Verification: 2 Burgers (₹150 ea) + 1 Pizza (₹250), Tier: 'GOLD', Coupon: 'FEAST50', Distance: 5.5 km -> Grand Total: ₹493.38 (Status: CONFIRMED)").runs[0].font.size = Pt(8.5)
    
    doc.add_paragraph("Streamlined Buggy Starter Code (Only 38 Lines | 10 Bugs Injected):").paragraph_format.space_after = Pt(1)
    add_code_box(doc, r3_streamlined_buggy, fill_hex="FFF1F2", border_hex="FECDD3")
    
    doc.add_paragraph("Identified Errors (10 Bugs Total):").paragraph_format.space_after = Pt(1)
    for b in r3_bugs_list:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1)
        clean_b = b.replace('<b>', '').replace('</b>', '').replace('<code>', '').replace('</code>', '')
        bp.add_run(clean_b).font.size = Pt(8)
        
    doc.add_paragraph("Verified Correct Solution:").paragraph_format.space_after = Pt(1)
    add_code_box(doc, r3_streamlined_fixed, fill_hex="F0FDF4", border_hex="BBF7D0")
    
    doc.save(output_path)
    print(f"Word document successfully updated at: {output_path}")

if __name__ == "__main__":
    out_docx = os.path.join("docs", "Code_Rescue_Round_Questions_and_Code.docx")
    generate_docx(out_docx)
