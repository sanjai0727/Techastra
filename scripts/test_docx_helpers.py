import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_box(doc, code_text, fill_hex="F8FAFC", border_hex="CBD5E1"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    lines = code_text.strip().split('\n')
    for i, line in enumerate(lines):
        line_num = f"{i+1:2d} | "
        r_num = p.add_run(line_num)
        r_num.font.name = 'Consolas'
        r_num.font.size = Pt(8)
        r_num.font.color.rgb = RGBColor(148, 163, 184)
        
        r_code = p.add_run(line + '\n' if i < len(lines)-1 else line)
        r_code.font.name = 'Consolas'
        r_code.font.size = Pt(8)
        r_code.font.color.rgb = RGBColor(15, 23, 42)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions ready")
